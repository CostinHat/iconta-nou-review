# banca_parser.py - Parser extras bancar XLS/XLSX/CSV (ING/Jasper).
# Logica validata pe ING real: grid Data/Detalii/Debit/Credit,
# semn dupa pozitia coloanei, randuri multi-linie grupate, formatul dupa primii octeti ai fisierului.
import re
from io import BytesIO

# [format_msg_v1] mesaj unic despre formatul asteptat, ridicat cand nu ies tranzactii.
FORMAT_ASTEPTAT = (
    "Formatul așteptat: coloanele Data, Detalii și fie Debit și Credit, fie o coloană Suma (cu semn). "
    "Separatorul de coloane: virgulă sau punct și virgulă; zecimalele cu punct sau virgulă."
)


def _celula_text(v, data_excel=False, datemode=0):
    """[lotul 12, deficiența 478 — „Extras .xlsx valid refuzat”] Celula unui tabel Excel, cum o citește parserul: o DATĂ (openpyxl dă
    `datetime` / `date`; xlrd dă numărul de serie Excel, cu tipul celulei „dată”) devine textul „zz.ll.aaaa”, ca în extrasele CSV. Înainte
    data rămânea obiect (`2026-10-16 00:00:00`), `_e_data` nu-l recunoștea, deci niciun rând nu era tranzacție: „Nu s-au găsit tranzacții
    în extras”. Celelalte valori trec neschimbate."""
    import datetime as _dt
    if data_excel and isinstance(v, (int, float)):
        import xlrd
        v = xlrd.xldate_as_datetime(v, datemode)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return "%02d.%02d.%04d" % (v.day, v.month, v.year)
    return v


#: [lotul 12, C121] un IBAN scris în antetul unui extras în tabel (deasupra rândului Data / Detalii / …): codul țării, 2 cifre de control
#: și restul alfanumeric — forma generală a IBAN-ului; spațiile dintre grupările de câte 4 sunt permise și se scot
_IBAN = re.compile(r"\b([A-Z]{2}[0-9]{2}(?: ?[A-Z0-9]){11,30})\b")


def iban_din_antet(randuri):
    """PURĂ. [C121] IBAN-ul contului din antetul tabelului (rândurile de DEASUPRA capului de tabel cu „Data”), sau None. Rândurile
    tranzacțiilor nu se citesc: acolo IBAN-ul e al partenerului, nu al contului extrasului. Fără cap de tabel, primele 10 rânduri."""
    for i, r in enumerate(randuri[:30]):
        jos = [str(c or "").lower() for c in r]
        if any("data" in c for c in jos) and any(("debit" in c or "suma" in c or "valoare" in c or "amount" in c) for c in jos):
            randuri = randuri[:i]
            break
    else:
        randuri = randuri[:10]
    for r in randuri:
        for c in r:
            m = _IBAN.search(str(c or "").upper())
            if m:
                v = m.group(1).replace(" ", "")
                # grupările cu spațiu pot înghiți cuvântul de după („… 0000 SOLD”): IBAN-ul românesc are 24 de caractere (RO + 2 + 4 + 16)
                return v[:24] if v.startswith("RO") else v
    return None


def _e_data(v):
    return bool(re.match(r"^\d{1,2}[./-]\d{1,2}[./-]\d{2,4}", str(v or "").strip()))

def _numar(v):
    s = str(v or "").strip().replace(" ", "")
    if not s:
        return None
    # ultimul separator (. sau ,) = zecimal; celelalte = mii
    up, uv = s.rfind("."), s.rfind(",")
    if up > uv:
        s = s.replace(",", "")
    elif uv > up:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def _detecteaza_delimitator(text):
    """[delim_robust_v1] Inlocuieste csv.Sniffer (care crapa cand ',' e si
    delimitator si zecimala, ex export RO '01.08.2026,Plata,100,50').
    Numara ; vs , vs TAB pe primele ~10 linii ne-goale; daca zecimalele par
    cu virgula prefera ';'; fallback ';' apoi ','."""
    linii = [l for l in text.splitlines() if l.strip()][:10]
    if not linii:
        return ","
    pv = sum(l.count(";") for l in linii)
    tb = sum(l.count("\t") for l in linii)
    vg = sum(l.count(",") for l in linii)
    # daca ';' apare, e delimitatorul (zecimalele RO folosesc virgula in acest caz)
    if pv > 0:
        return ";"
    if tb > 0 and tb >= vg:
        return "\t"
    if vg > 0:
        return ","
    return ";"


def _extrage_suma(r, celule, col, n_expected):
    """Suma dintr-o coloana unica 'Suma' (cu semn) sau din Debit/Credit separate.
    Cand ',' e delimitator SI zecimala, coloana Suma se sparge in mai multe celule
    decat coloane -> se reunesc surplusul (overflow rejoin)."""
    if col.get("suma") is not None:
        idx = col["suma"]
        if len(celule) > n_expected and idx == n_expected - 1:
            brut = ",".join(celule[idx:])
        else:
            brut = celule[idx] if idx < len(celule) else ""
        return _numar(brut) or 0
    deb = _numar(r[col["debit"]]) if col.get("debit") is not None and col["debit"] < len(r) else None
    cre = _numar(r[col["credit"]]) if col.get("credit") is not None and col["credit"] < len(r) else None
    return -(deb or 0) if deb else (cre or 0)


def _parse_grid(randuri):
    """randuri = lista de liste (celule). Detecteaza antet si extrage tranzactii.
    Antet acceptat: 'data' + (debit/credit SAU suma/valoare/amount).
    Sinonime detalii: detalii/descriere/explicatie."""
    antet_idx, col = None, {}
    for i, r in enumerate(randuri[:30]):
        jos = [str(c or "").lower() for c in r]
        are_data = any("data" in c for c in jos)
        are_debit = any("debit" in c for c in jos)
        are_suma = any(("suma" in c or "valoare" in c or "amount" in c) for c in jos)
        if are_data and (are_debit or are_suma):
            antet_idx = i
            for j, c in enumerate(jos):
                if "data" in c and "data" not in col: col["data"] = j
                if ("detali" in c or "descriere" in c or "explicat" in c
                        or "detaliu" in c) and "detalii" not in col:
                    col["detalii"] = j
                if "debit" in c and "debit" not in col: col["debit"] = j
                if "credit" in c and "credit" not in col: col["credit"] = j
                if ("suma" in c or "valoare" in c or "amount" in c) and "suma" not in col:
                    col["suma"] = j
            break
    if antet_idx is None:
        return []
    n_expected = max(col.values()) + 1 if col else 0
    tranzactii, curenta = [], None
    for r in randuri[antet_idx + 1:]:
        celule = [str(c or "").strip() for c in r]
        if not any(celule):
            continue
        data_v = celule[col["data"]] if col.get("data") is not None and col["data"] < len(celule) else ""
        if _e_data(data_v):
            if curenta: tranzactii.append(curenta)
            suma = _extrage_suma(r, celule, col, n_expected)
            det = celule[col["detalii"]] if col.get("detalii") is not None and col["detalii"] < len(celule) else ""
            curenta = {"data": data_v[:10], "detalii": det, "suma": round(suma, 2)}  # round-ok: suma din extras are deja bani; numai curățarea float
        elif curenta:
            extra = " ".join(c for c in celule if c)
            if extra: curenta["detalii"] = (curenta["detalii"] + " " + extra).strip()
    if curenta: tranzactii.append(curenta)
    return tranzactii


def _fara_antet(randuri):
    """[fara_antet_v1] Fallback fara antet: randuri care incep cu o data;
    ultima coloana numerica = suma (cu semn), restul intre data si suma = detalii."""
    tranzactii = []
    for r in randuri:
        celule = [str(c or "").strip() for c in r]
        if not celule or not _e_data(celule[0]):
            continue
        suma, suma_idx = None, None
        for j in range(len(celule) - 1, 0, -1):
            n = _numar(celule[j])
            if n is not None and celule[j] != "":
                suma, suma_idx = n, j
                break
        if suma is None:
            continue
        det = " ".join(c for c in celule[1:suma_idx] if c) if suma_idx else ""
        tranzactii.append({"data": celule[0][:10], "detalii": det, "suma": round(suma, 2)})  # round-ok: suma din extras are deja bani; numai curățarea float
    return tranzactii


def _sold_mt940(corp):
    """`:60F:` / `:62F:` (și `M`, soldul intermediar): `{C|D}AAMMZZ<valuta><suma>` -> `{suma, data, valuta}`; suma cu semn (D = sold
    debitor al contului clientului la bancă, adică negativ pentru firmă), data ISO (soldul de la sfârșitul acelei zile)."""
    m = re.match(r"([CD])(\d{6})([A-Z]{3})([\d,]+)", corp.strip())
    if not m:
        return None
    semn, aammzz, valuta, suma_s = m.groups()
    suma = float(suma_s.replace(",", "."))
    return {"suma": round(-suma if semn == "D" else suma, 2), "valuta": valuta,  # round-ok: suma din extras are deja bani; numai curățarea float
            "data": "%d-%s-%s" % (2000 + int(aammzz[:2]), aammzz[2:4], aammzz[4:6])}


def _parse_mt940(text):
    """Parser MT940 (SWIFT, format bancar standard). Sursa: specificatia SWIFT MT940 (Statement Message).
    `:25:` contul extrasului; `:60F:` / `:62F:` soldul initial / final; `:61:` linia (data, C/D, suma, tipul, referinta
    clientului[//referinta bancii]); `:86:` detaliile. Intoarce `{tranzactii, cont, sold_initial, sold_final}`; tranzactiile
    au aceeasi structura ca la `_parse_grid` ({data, detalii, suma}) plus `referinta`. `NONREF` = fara referinta -> "".
    [16.10.2026, lotul 6 pct.1, 277] Contul si referinta intra in cheia miscarii (`reconciliere_api.importa_extras`); soldurile se
    compara cu 5121 (276b)."""
    tranzactii, curenta = [], None
    cont = sold_initial = sold_final = None
    for linie_raw in text.splitlines():
        linie = linie_raw.strip()
        if linie.startswith(":25:") and cont is None:
            cont = linie[4:].strip() or None
        elif linie[:5] in (":60F:", ":60M:") and sold_initial is None:
            sold_initial = _sold_mt940(linie[5:])
        elif linie[:5] in (":62F:", ":62M:"):
            sold_final = _sold_mt940(linie[5:])
        elif linie.startswith(":61:"):
            if curenta:
                tranzactii.append(curenta)
            corp = linie[4:]
            # format :61: YYMMDD[MMDD]{C|D|RC|RD}[litera fondurilor]suma{N|F|S}<tip 3>referinta[//referinta bancii]
            m = re.match(r"(\d{6})(\d{4})?(R?[CD])([A-Z])?([\d,]+)(?:[NFS][A-Z0-9]{3})?(.*)", corp)
            if not m:
                curenta = None
                continue
            yymmdd, _mmdd, semn, _fond, suma_s, ref = m.groups()
            an = 2000 + int(yymmdd[:2])
            data = f"{yymmdd[4:6]}.{yymmdd[2:4]}.{an}"
            suma = float(suma_s.replace(",", "."))
            if semn in ("D", "RC"):  # D=debit(plata), RC=storno credit
                suma = -suma
            ref = (ref or "").strip()
            if ref.upper().startswith("NONREF"):
                ref = ref[6:].strip()
            curenta = {"data": data, "detalii": "", "suma": round(suma, 2), "referinta": ref}  # round-ok: suma din extras are deja bani; numai curățarea float
        elif linie.startswith(":86:") and curenta is not None:
            curenta["detalii"] = linie[4:].strip()
        elif curenta is not None and not linie.startswith(":") and linie:
            # continuare detalii pe linia urmatoare
            curenta["detalii"] = (curenta["detalii"] + " " + linie).strip()
    if curenta:
        tranzactii.append(curenta)
    return {"tranzactii": tranzactii, "cont": cont, "sold_initial": sold_initial, "sold_final": sold_final}


def citeste_extras(continut, nume_fisier=""):
    """Detecteaza formatul dupa primii octeti si intoarce `{tranzactii, cont, sold_initial, sold_final}`. Contul si soldurile vin
    din MT940; din tabele (xls / xlsx / csv) soldurile raman `None` — ecranul spune ca extrasul nu le contine —, iar contul e IBAN-ul
    din antetul tabelului, cand il are ([lotul 12, C121] `iban_din_antet`; comparat cu contul firmei inainte de import).
    Ridica ValueError (cu FORMAT_ASTEPTAT) daca parsarea esueaza sau da 0 tranzactii."""
    rez = {"cont": None, "sold_initial": None, "sold_final": None}
    try:
        if continut[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":  # OLE2 = .xls
            import xlrd
            wb = xlrd.open_workbook(file_contents=continut)
            sh = wb.sheet_by_index(0)
            randuri = [[_celula_text(sh.cell_value(r, c), sh.cell_type(r, c) == xlrd.XL_CELL_DATE, wb.datemode)
                        for c in range(sh.ncols)] for r in range(sh.nrows)]
            tranzactii = _parse_grid(randuri) or _fara_antet(randuri)
            rez["cont"] = iban_din_antet(randuri)
        elif continut[:2] == b"PK":  # zip = .xlsx
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(continut), read_only=True)
            sh = wb.active
            randuri = [[_celula_text(c) for c in rand] for rand in sh.iter_rows(values_only=True)]
            tranzactii = _parse_grid(randuri) or _fara_antet(randuri)
            rez["cont"] = iban_din_antet(randuri)
        else:  # text: MT940 sau csv
            # [lotul 12, clasa 251] codarea prin mecanismul unic al importurilor (BOM, UTF-8, Windows-1250 cu verificarea literelor);
            # înainte UTF-8 cu `errors="replace"`: diacriticele unui CSV Windows-1250 deveneau „�”
            from core import import_tabel as _it
            text = _it.decodeaza(continut, nume_fisier)[0] if isinstance(continut, bytes) else continut
            if ":61:" in text or ":20:" in text:  # marker MT940
                rez = _parse_mt940(text)
                tranzactii = rez.pop("tranzactii")
            else:
                import csv
                delim = _detecteaza_delimitator(text)
                randuri = list(csv.reader(text.splitlines(), delimiter=delim)) if text.strip() else []
                tranzactii = _parse_grid(randuri) or _fara_antet(randuri)
                rez["cont"] = iban_din_antet(randuri)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Nu s-a putut citi extrasul ({e}). {FORMAT_ASTEPTAT}")
    if not tranzactii:
        raise ValueError(f"Nu s-au găsit tranzacții în extras. {FORMAT_ASTEPTAT}")
    return dict(rez, tranzactii=tranzactii)


def parse_extras(continut, nume_fisier=""):
    """Numai tranzactiile din `citeste_extras` (ruta `parse-extras` si testele parserului)."""
    return citeste_extras(continut, nume_fisier)["tranzactii"]

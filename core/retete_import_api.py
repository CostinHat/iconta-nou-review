"""
core/retete_import_api.py - import retete HoReCa (F150) per firma.
Format: un rand per ingredient - Reteta; Pret vanzare (fara TVA); Ingredient; Cantitate/portie.
Gruparea pe reteta o face importul. Ingredientele se potrivesc pe articolele existente
dupa denumire (exact, apoi continere); nepotrivirile se raporteaza si reteta se sare.
Idempotenta: reteta existenta (aceeasi denumire) se sare cu motiv.
"""
from __future__ import annotations


def _gaseste_col(antet, *chei):
    for i, h in enumerate(antet):
        hl = str(h).strip().lower()
        if hl in chei:
            return i
    for i, h in enumerate(antet):
        hl = str(h).strip().lower()
        for k in chei:
            if k in hl:
                return i
    return -1


from core.numere import numar as _numar  # sursa unica (15.07.2026, vezi core/numere.py)


def extrage(continut, nume_fisier=""):
    """Intoarce [{denumire, pret, linii:[{ingredient, cantitate}]}] grupat pe reteta."""
    nume = (nume_fisier or "").lower()
    randuri = []
    if nume.endswith(".csv") or nume.endswith(".tsv") or nume.endswith(".txt"):
        import csv, io
        text = continut.decode("utf-8-sig", errors="replace") if isinstance(continut, bytes) else str(continut)
        prima = text.splitlines()[0] if text.splitlines() else ""
        delim = "\t" if nume.endswith(".tsv") else (";" if prima.count(";") > prima.count(",") else ",")
        randuri = list(csv.reader(io.StringIO(text), delimiter=delim))
    elif nume.endswith(".xlsx") or nume.endswith(".xlsm"):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(continut), read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            randuri.append(["" if c is None else c for c in r])
        wb.close()
    else:
        raise ValueError("format neacceptat (doar .csv sau .xlsx)")
    if not randuri:
        return []
    antet = [str(x) for x in randuri[0]]
    i_ret = _gaseste_col(antet, "reteta", "preparat", "produs finit", "denumire reteta")
    i_pret = _gaseste_col(antet, "pret", "pret vanzare")
    i_ing = _gaseste_col(antet, "ingredient", "articol", "materie")
    i_cant = _gaseste_col(antet, "cantitate", "cant", "consum")
    if i_ret < 0 or i_ing < 0 or i_cant < 0:
        raise ValueError("nu gasesc coloanele Reteta / Ingredient / Cantitate")
    grup = {}
    ordine = []
    for r in randuri[1:]:
        ret = str(r[i_ret] if i_ret < len(r) else "").strip()
        ing = str(r[i_ing] if i_ing < len(r) else "").strip()
        if not ret or not ing:
            continue
        cant = _numar(r[i_cant]) if 0 <= i_cant < len(r) else 0.0
        pret = _numar(r[i_pret]) if 0 <= i_pret < len(r) else 0.0
        if ret not in grup:
            grup[ret] = {"denumire": ret[:255], "pret": pret, "linii": []}
            ordine.append(ret)
        if pret > 0:
            grup[ret]["pret"] = pret
        grup[ret]["linii"].append({"ingredient": ing[:255], "cantitate": cant})
    return [grup[k] for k in ordine]


def potriveste(conn, schema, retete):
    """Potriveste ingredientele pe articole. Adauga articol_id sau motiv per linie
    + valid/motiv per reteta (toate liniile potrivite si cantitati > 0)."""
    with conn.cursor() as cur:
        cur.execute(f"SELECT id, denumire FROM {schema}.articole")
        arts = cur.fetchall()
    exact = {d.strip().lower(): i for i, d in arts}
    for ret in retete:
        motive = []
        for l in ret["linii"]:
            cheie = l["ingredient"].strip().lower()
            aid = exact.get(cheie)
            if aid is None:
                cand = [i for i, d in arts if cheie in d.strip().lower() or d.strip().lower() in cheie]
                aid = cand[0] if len(cand) == 1 else None
            l["articol_id"] = aid
            if aid is None:
                motive.append(f"ingredient negasit: {l['ingredient']}")
            elif l["cantitate"] <= 0:
                motive.append(f"cantitate invalida la {l['ingredient']}")
        ret["valid"] = not motive
        ret["motiv"] = "; ".join(motive) if motive else "ok"
    return retete


def rezumat(retete):
    ok = [r for r in retete if r.get("valid")]
    return {"total": len(retete), "valide": len(ok), "invalide": len(retete) - len(ok),
            "ingrediente": sum(len(r["linii"]) for r in retete)}


def importa(conn, schema, retete):
    """Creeaza retetele valide prin retete_api.salveaza. Existentele/invalidele se sar."""
    from core import retete_api
    create, sarite = 0, []
    with conn.cursor() as cur:
        for ret in retete:
            if not ret.get("valid"):
                sarite.append({"denumire": ret["denumire"], "motiv": ret.get("motiv", "invalid")})
                continue
            cur.execute(f"SELECT id FROM {schema}.retete WHERE lower(denumire)=lower(%s)", (ret["denumire"],))
            if cur.fetchone():
                sarite.append({"denumire": ret["denumire"], "motiv": "exista deja"})
                continue
            retete_api.salveaza(conn, schema, {"denumire": ret["denumire"], "pret_fara_tva": ret.get("pret") or 0,
                "linii": [{"articol_id": l["articol_id"], "cantitate": l["cantitate"]} for l in ret["linii"]]})
            create += 1
    return {"create": create, "sarite": sarite}

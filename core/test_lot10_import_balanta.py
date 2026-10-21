# -*- coding: utf-8 -*-
"""[lotul 10, deficiența 425 — PRIORITAR, comanda Costin 10.1, verbatim în DECIZII „Comanda 10.1”] Importul balanței de deschidere.

Verbatim: „O balanță completă la 30.09.2026 (format SAGA: solduri inițiale an, rulaje precedente, rulaje lună, total sume, solduri finale,
antet pe două rânduri) e citită doar din «Solduri finale». Soldurile de la 1 ianuarie și rulajele ianuarie–septembrie se pierd, iar
conturile cu sold final 0 (707, 607, 641, 646, 6811, 4426, 4427) nu intră deloc. […] Când fișierul are soldurile inițiale ale anului și
rulajele, ele intră ca sold inițial și sume precedente, cu toate conturile care au rulaj. Când are doar solduri finale, comportamentul de
acum rămâne, dar aplicația spune clar că rulajele anului lipsesc și ce afectează. Tot aici: când balanța are sold pe 371, metoda de stoc
se cere la migrare, nu rămâne «nedeclarată».”

TEMEI (citate verbatim):
- manualul SAGA C, „Preluare date contabile” (anaf_surse/saga_manual_topic12_preluare_date_contabile.txt): coloanele „Debit precedent” /
  „Credit precedent” se completează „cu rulajele cumulate (inclusiv soldurile iniţiale) - coloanele «Total Sume debitoare/creditoare»
  din balanţa lunii anterioare” -> TOTAL SUME = sold inițial an + rulajele anului; rulajele = total sume − sold inițial;
- OMFP 2634/2015 anexa 1 pct.52: „Soldul debitor sau creditor al fiecărui cont se stabilește în funcție de rulajele debitoare și
  creditoare ale contului respectiv, ținându-se seama de soldul de la începutul anului” -> sold inițial + rulaje = sold final, pe cont;
- OMFP 1802/2014 pct.594 (planul de conturi): „371. Mărfuri (A)”, „378. Diferențe de preț la mărfuri (A/P)”.

Balanța de probă e INVENTATĂ (date sintetice, nicio firmă reală), echilibrată pe fiecare pereche: 1 ian. 2026 capital 30.000, rezultat
reportat 5.000, echipamente 20.000, bancă 12.000, mărfuri 3.000; în ian.–sep.: vânzări 180.000 + TVA 21% 37.800, încasări 200.000,
cumpărări de mărfuri 120.000 + TVA 25.200, plăți 140.000, descărcarea mărfii 110.000, salarii 30.000, contribuția asiguratorie 675,
amortizare 2.000, regularizarea și plata TVA 12.600, închiderea lunară a veniturilor și cheltuielilor pe 121. Septembrie = 1/5 din
fiecare operație (sumele precedente = total sume − rulajele lunii). Fișierul `date_test/import_saga/balanta_saga_completa.xlsx` e
scris de `randuri_saga()` (e2e-ul îl încarcă).
"""
import contextlib
import os
import re

import pytest

from core import db as _db, import_tabel as it, solduri_api

RAD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FISIER = os.path.join(RAD, "date_test", "import_saga", "balanta_saga_completa.xlsx")

#: (cont, denumire, sold inițial debit, sold inițial credit, rulaj debit, rulaj credit) — rulajele 01.01–30.09.2026, fără soldul inițial
BALANTA = [
    ("1012", "Capital subscris vărsat", 0, 30000, 0, 0),
    ("117", "Rezultatul reportat", 0, 5000, 0, 0),
    ("121", "Profit și pierdere", 0, 0, 142675, 180000),
    ("2131", "Echipamente tehnologice", 20000, 0, 0, 0),
    ("2813", "Amortizarea instalațiilor", 0, 0, 0, 2000),
    ("371", "Mărfuri", 3000, 0, 120000, 110000),
    ("401", "Furnizori", 0, 0, 140000, 145200),
    ("4111", "Clienți", 0, 0, 217800, 200000),
    ("421", "Personal - salarii datorate", 0, 0, 30000, 30000),
    ("436", "Contribuția asiguratorie pentru muncă", 0, 0, 675, 675),
    ("4423", "TVA de plată", 0, 0, 12600, 12600),
    ("4426", "TVA deductibilă", 0, 0, 25200, 25200),
    ("4427", "TVA colectată", 0, 0, 37800, 37800),
    ("5121", "Conturi la bănci în lei", 12000, 0, 200000, 183275),
    ("607", "Cheltuieli privind mărfurile", 0, 0, 110000, 110000),
    ("641", "Cheltuieli cu salariile personalului", 0, 0, 30000, 30000),
    ("646", "Cheltuieli privind contribuția asiguratorie pentru muncă", 0, 0, 675, 675),
    ("6811", "Cheltuieli de exploatare privind amortizarea imobilizărilor", 0, 0, 2000, 2000),
    ("707", "Venituri din vânzarea mărfurilor", 0, 0, 180000, 180000),
]


def _sume(mesaj):
    """Sumele numite într-un refuz, în ordine (format `bani()`: 1.234,56 lei)."""
    return re.findall(r"(-?\d{1,3}(?:\.\d{3})*,\d{2}) lei", str(mesaj))


def _sf(si_d, si_c, r_d, r_c):
    net = si_d - si_c + r_d - r_c
    return (net, 0) if net > 0 else (0, -net)


def randuri_saga():
    """Balanța de verificare în forma exportului SAGA: titluri, antet pe DOUĂ rânduri (grupul deasupra, Debit / Credit dedesubt),
    cele cinci perechi, rândul de total."""
    out = [["Balanță de verificare analitică"], ["SC Firmă Probă Rulaje SRL"], ["Perioada: 01.09.2026 - 30.09.2026"],
           ["Simbol cont", "Denumirea contului", "Solduri inițiale an", "", "Sume precedente", "", "Rulaje perioadă", "",
            "Total sume", "", "Solduri finale", ""],
           ["", "", "Debit", "Credit", "Debit", "Credit", "Debit", "Credit", "Debit", "Credit", "Debit", "Credit"]]
    tot = [0] * 10
    for cont, den, si_d, si_c, r_d, r_c in BALANTA:
        ts = (si_d + r_d, si_c + r_c)
        rl = (r_d // 5, r_c // 5)
        sp = (ts[0] - rl[0], ts[1] - rl[1])
        sf = _sf(si_d, si_c, r_d, r_c)
        vals = [si_d, si_c, sp[0], sp[1], rl[0], rl[1], ts[0], ts[1], sf[0], sf[1]]
        tot = [a + b for a, b in zip(tot, vals)]
        out.append([cont, den] + vals)
    out.append(["TOTAL", ""] + tot)
    return out


def _csv(randuri):
    return "\n".join(";".join(str(c) for c in r) for r in randuri).encode("utf-8")


def _pe_cont(d):
    return {r["cont"]: r for r in d["randuri"]}


# ── citirea ─────────────────────────────────────────────────────────────────────────────────────────────────────────────────
def test_fixtura_saga_e_echilibrata_pe_fiecare_pereche():
    """Fixtura însăși: fiecare pereche are debit = credit (altfel testele de mai jos ar măsura o balanță imposibilă)."""
    tot = randuri_saga()[-1][2:]
    assert [tot[i] == tot[i + 1] for i in range(0, 10, 2)] == [True] * 5, tot
    assert tot[-2:] == [79525, 79525]


def test_def_425_balanta_completa_intra_cu_soldul_initial_si_rulajele():
    """ROȘU înainte: numai soldurile finale, iar 707 / 607 / 641 / 646 / 6811 / 4426 / 4427 (sold final 0) lipseau (fara_sold = 10).
    MUTAȚIA: `perechi_balanta` care întoarce {} -> rulajele lipsesc și conturile cu sold final 0 nu intră -> pică."""
    d = solduri_api.extrage_balanta_detaliat(open(FISIER, "rb").read(), "balanta_saga_completa.xlsx")
    pc = _pe_cont(d)
    assert d["rulaje"] is True and d["fara_sold"] == 0 and d["ignorate"] == ["TOTAL"]
    assert set(pc) == {b[0] for b in BALANTA}            # „cu toate conturile care au rulaj” (comanda 10.1)
    # cifra de afaceri a anului: 707 = 180.000 lei, cu sold final 0 (închis pe 121) — SAGA topic12: total sume = sold inițial + rulaje
    assert (pc["707"]["debit"], pc["707"]["credit"], pc["707"]["rulaj_debit"], pc["707"]["rulaj_credit"]) == (0.0, 0.0, 180000.0, 180000.0)
    for cont in ("607", "641", "646", "6811", "4426", "4427"):
        assert pc[cont]["debit"] == pc[cont]["credit"] == 0.0 and pc[cont]["rulaj_debit"] > 0, cont
    # 5121: sold inițial 12.000 + 200.000 − 183.275 = 28.725 (OMFP 2634/2015 anexa 1 pct.52: soldul din rulaje „ținându-se seama de
    # soldul de la începutul anului”)
    assert (pc["5121"]["si_debit"], pc["5121"]["rulaj_debit"], pc["5121"]["rulaj_credit"], pc["5121"]["debit"]) == (
        12000.0, 200000.0, 183275.0, 28725.0)
    assert d["perechi"]["si"] == ["Solduri inițiale an Debit", "Solduri inițiale an Credit"]
    assert set(d["perechi"]) == {"si", "sp", "rl", "ts"}


def test_425_fara_total_sume_rulajele_vin_din_precedente_plus_luna():
    """Fără „Total sume”: total sume = sume precedente + rulajele lunii (aceeași balanță). MUTAȚIA: ramura `ts or din_parti` -> pică."""
    rr = [r[:8] + r[10:] for r in randuri_saga()]          # scoate perechea „Total sume”
    pc = _pe_cont(solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv"))
    assert (pc["707"]["rulaj_debit"], pc["707"]["rulaj_credit"]) == (180000.0, 180000.0)
    assert (pc["371"]["si_debit"], pc["371"]["rulaj_debit"], pc["371"]["rulaj_credit"]) == (3000.0, 120000.0, 110000.0)


def test_425_balanta_care_se_contrazice_se_refuza_numit():
    """Soldul final al lui 5121 schimbat (28.725 -> 28.000): sold inițial + rulaje nu mai dau soldul final. MUTAȚIA: verificarea
    ecuației scoasă din `verifica_rulaje` -> nu mai pică (și nici echilibrul, fiindcă totalul de jos nu se recitește)."""
    rr = randuri_saga()
    for r in rr:
        if r[0] == "5121":
            r[10] = 28000
        if r[0] == "1012":
            r[11] = 29275                                     # soldurile finale rămân echilibrate: numai ecuația pe cont pică
    with pytest.raises(ValueError) as e:
        solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv")
    # structura refuzului: rândul și contul numite, cele două cifre ale ecuației (sold inițial + rulaje vs. sold final, net) — nu „apare
    # undeva «contul 1012»”. OMFP 2634/2015 anexa 1 pct.52: „Soldul debitor sau creditor al fiecărui cont se stabilește în funcție de
    # rulajele debitoare și creditoare ale contului respectiv, ținându-se seama de soldul de la începutul anului”
    # (locul și cifrele, nu formularea: o rescriere a frazei după DS cap.6 nu sparge proba)
    m = re.match(r"Rândul (\d+) \(contul (\w+)\): ", str(e.value))
    assert m and m.groups() == ("6", "1012") and _sume(e.value) == ["-30.000,00", "-29.275,00"], str(e.value)


def test_425_precedente_plus_luna_diferit_de_total_sume_se_refuza():
    rr = randuri_saga()
    for r in rr:
        if r[0] == "707":
            r[7] += 100                                       # rulajul lunii, credit
    with pytest.raises(it.FisierNecitit) as e:
        solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv")
    # rândul, contul și cele patru cifre: precedente + luna (debit, credit) față de „total sume” (debit, credit) — manualul SAGA C:
    # „Debit precedent” / „Credit precedent” = „rulajele cumulate (inclusiv soldurile iniţiale)”, deci precedente + lună = total sume
    m = re.match(r"Rândul (\d+) \(contul (\w+)\): ", str(e.value))
    assert m and m.groups() == ("24", "707") and _sume(e.value) == ["180.000,00", "180.100,00", "180.000,00", "180.000,00"], str(e.value)


def test_425_jumatate_de_pereche_se_refuza_numit():
    rr = [r[:3] + r[4:] for r in randuri_saga()]              # scoate „Solduri inițiale an Credit”
    with pytest.raises(it.FisierNecitit) as e:
        solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv")
    # grupul numit și coloanele lui găsite (una singură, jumătatea de pereche), exact
    assert re.findall(r"„([^”]+)”", str(e.value)) == ["solduri inițiale an"], str(e.value)
    assert re.findall(r"«([^»]+)»", str(e.value)) == ["b.csv", "Solduri inițiale an Debit"], str(e.value)   # fișierul, apoi coloanele


def test_425_numai_solduri_finale_comportamentul_ramane_si_lipsa_se_spune():
    """Comanda 10.1: „Când are doar solduri finale, comportamentul de acum rămâne, dar aplicația spune clar că rulajele anului lipsesc
    și ce afectează.” MUTAȚIA: textul fără „D101” -> pică."""
    rr = [r[:2] + r[10:] for r in randuri_saga()]
    d = solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv")
    assert d["rulaje"] is False and d["fara_sold"] == 10 and "rulaj_debit" not in d["randuri"][0]
    t = solduri_api.avertisment_fara_rulaje("2026-09-30")
    for ce in ("până la 30.09.2026 lipsesc", "(D101)", "situațiile financiare anuale", "plafonul de microîntreprindere",
               "controalele cumulate de TVA", "„sold inițial” / „sume precedente”"):
        assert ce in t, ce
    assert solduri_api.avertisment_fara_rulaje("2025-12-31") is None    # la 31 decembrie nu lipsește nimic din anul preluării


def test_425_verifica_rulaje_cere_toate_randurile_si_echilibrul():
    r = [{"cont": "5121", "debit": 100, "credit": 0, "si_debit": 100, "si_credit": 0, "rulaj_debit": 0, "rulaj_credit": 0},
         {"cont": "1012", "debit": 0, "credit": 100}]
    with pytest.raises(ValueError) as e:
        solduri_api.verifica_rulaje(r)
    assert str(e.value).startswith("Balanța are soldurile inițiale ale anului și rulajele numai pe o parte din conturi")
    r[1].update({"si_debit": 0, "si_credit": 90, "rulaj_debit": 0, "rulaj_credit": 10})
    with pytest.raises(ValueError) as e:
        solduri_api.verifica_rulaje(r)
    assert str(e.value).startswith("Soldurile inițiale ale anului nu se echilibrează: debit 100,00 lei, credit 90,00 lei")


def test_425_sold_pe_marfa_cere_metoda():
    assert solduri_api.are_sold_marfa([{"cont": "371.01", "debit": 5, "credit": 0}])
    assert solduri_api.are_sold_marfa([{"cont": "378", "debit": 0, "credit": 5}])
    assert not solduri_api.are_sold_marfa([{"cont": "371", "debit": 0, "credit": 0}, {"cont": "3021", "debit": 5, "credit": 0}])


# ── scrierea (schemă efemeră din template, tranzacție anulată) ─────────────────────────────────────────────────────────────────
SCHEMA = "ztest_lot10_balanta"


def _db_ok():
    try:
        _db.init_pool()
        with _db.get_conn():
            return True
    except Exception:
        return False


@pytest.fixture
def conn():
    if not _db_ok():
        pytest.skip("DB indisponibil")
    from core import tenant_provisioning as _tp
    with _db.get_conn() as c:
        try:
            with c.cursor() as cur:
                cur.execute("DROP SCHEMA IF EXISTS %s CASCADE" % SCHEMA)
                cur.execute(_tp.parametrizeaza_template(open(os.path.join(RAD, "tenant_template.sql"), encoding="utf-8").read(), SCHEMA))
                cur.execute("SET LOCAL search_path TO %s, public" % SCHEMA)
                cur.execute("INSERT INTO firma_profil (id, nume, cui) VALUES (1, 'FIRMĂ PROBĂ RULAJE SRL', '14399840')")
            yield c
        finally:
            c.rollback()


def _uc(monkeypatch, c):
    """Use-case-ul pe conexiunea testului (tranzacția anulată la final), fără firmă în portofoliu."""
    from core import uc_tenants, uc_comun

    @contextlib.contextmanager
    def _gc(schema=None):
        yield c
    monkeypatch.setattr(uc_tenants.db, "get_conn", _gc)
    monkeypatch.setattr(uc_comun, "_schema_sau_404", lambda ctx, tid: SCHEMA)
    monkeypatch.setattr(uc_tenants, "_data_soldurilor", lambda conn, schema: "2026-09-30")
    return uc_tenants


def test_def_425_importa_scrie_soldul_initial_si_rulajele(conn):
    """MUTAȚIA: INSERT-ul fără rulaje (None în loc de valori) -> 707 fără rulaj în bază -> pică."""
    d = solduri_api.extrage_balanta_detaliat(open(FISIER, "rb").read(), "b.xlsx")
    rez = solduri_api.importa(conn, d["randuri"], "2026-09-30")
    assert (rez["randuri"], rez["rulaje"], rez["avertisment_rulaje"]) == (19, True, None)
    with conn.cursor() as cur:
        cur.execute("SELECT cont, sold_debitor, sold_creditor, si_an_debit, si_an_credit, rulaj_an_debit, rulaj_an_credit "
                    "FROM solduri_initiale WHERE cont IN ('707', '5121') ORDER BY cont")
        assert [tuple(float(x) if x is not None and not isinstance(x, str) else x for x in r) for r in cur.fetchall()] == [
            ("5121", 28725.0, 0.0, 12000.0, 0.0, 200000.0, 183275.0),
            ("707", 0.0, 0.0, 0.0, 0.0, 180000.0, 180000.0)]       # cifra de afaceri a anului: 180.000 lei
        cur.execute("SELECT sum(rulaj_an_debit), sum(rulaj_an_credit), sum(si_an_debit), sum(si_an_credit) FROM solduri_initiale")
        assert [float(x) for x in cur.fetchone()] == [1249425.0, 1249425.0, 35000.0, 35000.0]


def test_425_numai_solduri_finale_se_scrie_ca_pana_acum_cu_avertisment(conn):
    rr = [r[:2] + r[10:] for r in randuri_saga()]
    d = solduri_api.extrage_balanta_detaliat(_csv(rr), "b.csv")
    rez = solduri_api.importa(conn, d["randuri"], "2026-09-30")
    # avertismentul importului e EXACT al regulii pentru data balanței (conținutul lui — D101 etc. — e păzit de testul de mai sus)
    assert rez["rulaje"] is False and rez["avertisment_rulaje"] == solduri_api.avertisment_fara_rulaje("2026-09-30") is not None
    with conn.cursor() as cur:
        cur.execute("SELECT count(*), count(rulaj_an_debit) FROM solduri_initiale")
        assert cur.fetchone() == (9, 0)                     # 9 conturi cu sold final; rulajele NULL = „numai solduri finale”


def test_def_425_metoda_de_stoc_se_cere_la_migrare(conn, monkeypatch):
    """Sold pe 371 + firmă fără metodă -> salvarea fără alegere se refuză NUMIT (camp metoda_stoc), nimic scris; cu alegerea, metoda intră
    în Date firmă și în jurnalul lor. MUTAȚIA: `cere_metoda` forțat False -> salvarea trece fără metodă -> pică."""
    import main
    from core import erori, metoda_stoc
    uc = _uc(monkeypatch, conn)
    d = solduri_api.extrage_balanta_detaliat(open(FISIER, "rb").read(), "b.xlsx")
    corp = main.SolduriIn(randuri=[main.SoldRand(**{k: v for k, v in r.items() if k in main.SoldRand.model_fields}) for r in d["randuri"]],
                          data_referinta="2026-09-30")
    with pytest.raises(erori.DateInvalide) as e:
        uc.solduri_salveaza(1, corp, {"uid": 7})
    assert e.value.args[0]["erori_campuri"][0]["camp"] == "metoda_stoc"
    with conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM solduri_initiale")
        assert cur.fetchone()[0] == 0
    corp.metoda_stoc = metoda_stoc.CV
    rez = uc.solduri_salveaza(1, corp, {"uid": 7})
    assert rez["metoda_stoc"] == metoda_stoc.CV and rez["rulaje"] is True
    with conn.cursor() as cur:
        cur.execute("SELECT metoda_stoc FROM firma_profil")
        assert cur.fetchone()[0] == "cantitativ_valoric"
        cur.execute("SELECT camp, valoare_veche, valoare_noua, user_id FROM firma_profil_jurnal")
        assert cur.fetchall() == [("metoda_stoc", None, "cantitativ_valoric", 7)]


def test_425_metoda_la_migrare_aceleasi_reguli_ca_date_firma(conn):
    from core import metoda_stoc as ms
    with conn.cursor() as cur:
        with pytest.raises(ValueError) as e:
            ms.declara_la_migrare(cur, SCHEMA, "cantitativ_valoric_fifo", 7)
        assert e.value.cod == ms.COD_NESUPORTATA
        with pytest.raises(ValueError):
            ms.declara_la_migrare(cur, SCHEMA, "fifo", 7)
        assert ms.declara_la_migrare(cur, SCHEMA, ms.GV, 7) == ms.GV
        assert ms.declara_la_migrare(cur, SCHEMA, ms.CV, 7) == ms.GV          # declarată: nu se schimbă de aici (Date firmă, cu jurnal)


def test_425_previzualizarea_spune_rulajele_si_cere_metoda(conn, monkeypatch):
    uc = _uc(monkeypatch, conn)
    import json
    r = json.loads(uc.solduri_incarca(1, open(FISIER, "rb").read(), "balanta_saga_completa.xlsx", {"uid": 7}).body)
    assert (r["rulaje"], r["avertisment_rulaje"], r["cere_metoda_stoc"], len(r["metode_stoc"]), r["existenta"]) == (True, None, True, 2, None)
    rr = [x[:2] + x[10:] for x in randuri_saga()]
    r = json.loads(uc.solduri_incarca(1, _csv(rr), "b.csv", {"uid": 7}).body)
    assert r["rulaje"] is False and r["avertisment_rulaje"].startswith("Fișierul are numai soldurile finale")


def test_scriere_fixtura_e_la_zi():
    """Fișierul din date_test e exact `randuri_saga()` (generat de el) — e2e-ul încarcă aceleași cifre."""
    from openpyxl import load_workbook
    wb = load_workbook(FISIER, read_only=True)
    try:
        citit = [["" if c is None else c for c in r] for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()
    asteptat = [list(r) + [""] * (12 - len(r)) for r in randuri_saga()]
    assert citit == asteptat

"""
core/istoric_declaratii_import_api.py — import istoric declaratii (stratul 7) per firma.

Citeste o lista de declaratii deja depuse (.xlsx/.csv): tip, an, luna, data depunere.
Scrie in tabelul global public.declaratii_depuse cu sursa='migrare' (ca sa nu apara fals
restanta). Import per firma: DELETE WHERE tenant_id=X AND sursa='migrare', apoi INSERT.
NU sterge declaratiile depuse prin iConta (sursa='iconta').
"""
from __future__ import annotations
import datetime

from core import afirmatii as _af
from core.unde import Unde as _Unde  # [P8] fiecare avertisment stie pe ce rand e
from core.migrare_api import respinge  # [P8/C] respingerea e o afirmatie, cu regula din nomenclator

# tipuri cunoscute (pentru normalizare + recunoastere)
_TIPURI = ["D100", "D101", "D112", "D205", "D300", "D301", "D390", "D394", "D406"]


def _gaseste_col(antet, *chei):
    for i, h in enumerate(antet):
        hl = str(h).strip().lower()
        for k in chei:
            if k in hl:
                return i
    return -1


def _intreg(v):
    try:
        return int(round(float(str(v).strip().replace(",", ".")))) if str(v).strip() else 0
    except ValueError:
        return 0


def _normalizeaza_tip(v):
    """'100' / 'd100' / 'Declaratia 100' -> 'D100'. Necunoscut -> uppercase trim."""
    t = str(v or "").strip().upper().replace(" ", "")
    # extrage cifrele
    cifre = "".join(ch for ch in t if ch.isdigit())
    if cifre:
        cand = "D" + cifre
        if cand in _TIPURI:
            return cand, True
    if t in _TIPURI:
        return t, True
    # poate are deja Dxxx
    for tip in _TIPURI:
        if tip in t:
            return tip, True
    return (t or "?"), False   # necunoscut


def _data(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    t = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(t, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def extrage(continut, nume_fisier=""):
    """Intoarce [{tip, an, luna, data_depunere, tip_cunoscut, ok, avertisment}]."""
    nume = (nume_fisier or "").lower()
    randuri = []

    if nume.endswith(".csv") or nume.endswith(".tsv") or nume.endswith(".txt"):
        import csv, io
        text = continut.decode("utf-8-sig", errors="replace") if isinstance(continut, bytes) else str(continut)
        prima = text.splitlines()[0] if text.splitlines() else ""
        delim = "\t" if nume.endswith(".tsv") else (";" if prima.count(";") > prima.count(",") else ",")
        randuri = list(csv.reader(io.StringIO(text), delimiter=delim))
    elif nume.endswith(".xlsx") or nume.endswith(".xlsm"):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(continut), read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            randuri.append(["" if c is None else c for c in r])
        wb.close()
    else:
        raise ValueError("format neacceptat (doar .csv sau .xlsx)")

    if not randuri:
        return []

    antet = [str(x) for x in randuri[0]]
    i_tip = _gaseste_col(antet, "tip", "declaratie", "declarație", "formular")
    i_an = _gaseste_col(antet, "an")
    i_luna = _gaseste_col(antet, "luna", "lună", "perioada", "perioadă")
    i_data = _gaseste_col(antet, "depunere", "depus", "data")
    if i_tip < 0:
        raise ValueError("nu găsesc coloana cu tipul declarației (tip/declarație/formular) - fișier nerecunoscut")

    out = []
    for i, r in enumerate(randuri[1:], start=2):
        def cel(i):
            return str(r[i]).strip() if (0 <= i < len(r)) else ""
        tip_raw = cel(i_tip)
        if not tip_raw:
            continue
        tip, cunoscut = _normalizeaza_tip(tip_raw)
        an = _intreg(cel(i_an))
        luna = _intreg(cel(i_luna))
        data_dep = _data(r[i_data]) if (0 <= i_data < len(r)) else None

        # [P8, 22.08] fiecare avertisment e AFIRMATIA lui, nu un sir intr-o lista. O lista de texte
        # nu poate fi numarata pe regula, nu poate fi filtrata si nu spune despre CE rand vorbeste -
        # exact forma pe care campania o desfiinteaza. Randorul citeste `.motiv`.
        _unde = _Unde("rand", i)
        avert = []
        if not cunoscut:
            avert.append(_af.afirmatie("neconformitate", "declarație depusă",
                                       "tip necunoscut: %s" % (tip or "?"),
                                       unde=_unde, regula="tip_necunoscut"))
        if not (2018 <= an <= 2027):
            avert.append(_af.afirmatie("neconformitate", "declarație depusă",
                                       "an neplauzibil: %s" % an,
                                       unde=_unde, regula="an_invalid"))
        if not (1 <= luna <= 12):
            avert.append(_af.afirmatie("neconformitate", "declarație depusă",
                                       "lună invalidă: %s" % luna,
                                       unde=_unde, regula="luna_invalida"))

        out.append({"tip": tip, "an": an, "luna": luna, "data_depunere": data_dep,
                    "tip_cunoscut": cunoscut, "avertisment": avert, "ok": len(avert) == 0})
    return out


def rezumat(conn, tenant_id):
    """{are_istoric, randuri} pentru o firma (doar randurile din migrare).
    Coloana `sursa` e garantata de migrare_declaratii_depuse_randuri (public), nu mai lazy."""
    with conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM public.declaratii_depuse WHERE tenant_id=%s AND sursa='migrare'",
                    (tenant_id,))
        n = cur.fetchone()[0]
    return {"are_istoric": n > 0, "randuri": n}


TIPURI_CUNOSCUTE = ("D100", "D101", "D112", "D205", "D300", "D301", "D390", "D394",
                    "D406", "D212", "S1003", "S1005", "D394A", "D311", "D207")


def verifica_randuri(randuri, azi=None):
    """PURA: [{rand, motiv, mesaj}]. Nicio validare pana la 15.07.2026: tip D999
    inexistent, luna 13, data depunerii inaintea perioadei raportate - toate intrau."""
    import datetime as _dt
    azi = azi or _dt.date.today()
    er = []
    for i, r in enumerate(randuri or [], start=2):
        tip = str(r.get("tip") or "").strip().upper()   # VALIDARE contra TIPURI_CUNOSCUTE (uppercase, nomenclator ANAF)
        an = int(r.get("an") or 0)
        luna = r.get("luna")
        if tip and tip not in TIPURI_CUNOSCUTE:
            er.append(respinge("declarație depusă", i, "tip_necunoscut",
                               "declarația %s nu există în nomenclatorul ANAF" % tip))
        if not (2000 <= an <= azi.year + 1):
            er.append(respinge("declarație depusă", i, "an_invalid",
                               "%s: anul %s e în afara intervalului" % (tip or "?", an)))
        if luna is not None and str(luna).strip() != "":
            try:
                l = int(luna)
                if not (1 <= l <= 12):
                    er.append(respinge("declarație depusă", i, "luna_invalida",
                                       "%s: luna %s (așteptat 1-12)" % (tip or "?", l)))
            except (TypeError, ValueError):
                er.append(respinge("declarație depusă", i, "luna_invalida",
                                   "%s: luna %r nu e număr" % (tip or "?", luna)))
        d = r.get("data_depunere")
        if d:
            try:
                dd = d if isinstance(d, _dt.date) else _dt.date.fromisoformat(str(d)[:10])
                l = int(luna) if str(luna or "").strip().isdigit() else 12
                sfarsit = _dt.date(an, l, 1) if 2000 <= an <= 2100 and 1 <= l <= 12 else None
                if sfarsit and dd < sfarsit:
                    er.append(respinge("declarație depusă", i, "data_inainte_de_perioada",
                                       "%s %s/%s: depusă la %s, înainte de perioada raportată"
                                       % (tip or "?", luna, an, dd)))
            except ValueError:
                er.append(respinge("declarație depusă", i, "data_invalida",
                                   "%s: data depunerii %r nu se înțelege" % (tip or "?", d)))
    return er


def importa(conn, tenant_id, randuri):
    """DELETE doar randurile de migrare ale firmei + INSERT. NU atinge sursa='iconta'.
    Ridica ValueError daca randurile nu pot intra (vezi verifica_randuri)."""
    # [probare invalid, 03.09.2026] UN IMPORT GOL STERGEA ISTORICUL, IN TACERE.
    # `randuri=[]` trecea de verificare (n-are ce respinge), ajungea la DELETE-ul de mai jos, care
    # sterge TOATE randurile de migrare ale firmei, nu insera nimic, si raspundea `200 {"importati": 0}`.
    # Adica exact forma cea mai rea: o pierdere de date raportata ca succes. Un fisier gol, o coloana
    # necitita, un filtru care n-a potrivit nimic — oricare din ele ajungea aici.
    # Istoricul declaratiilor depuse sta la baza termenelor si a controlului fiscal.
    if not randuri:
        raise ValueError("nu ai trimis niciun rând. Importul ar fi șters istoricul de declarații "
                         "încărcat până acum pentru firma asta și n-ar fi pus nimic în loc. "
                         "Dacă chiar vrei să golești istoricul importat, e altă operațiune.")
    er = verifica_randuri(randuri)
    if er:
        # [probare invalid, 03.09.2026] Se numără RÂNDURILE, nu erorile: un singur rând gol
        # produce două erori, iar mesajul spunea „2 rânduri nu pot intra" despre un rând. O
        # cifră falsă într-un refuz e mai rea decât un refuz sec — omul caută al doilea rând.
        _nr = len({x["rand"] for x in er})
        raise ValueError("%s nu %s intra: %s. Istoricul declarațiilor stă la "
                         "baza termenelor și a controlului fiscal."
                         % ("Un rând" if _nr == 1 else "%d rânduri" % _nr,
                            "poate" if _nr == 1 else "pot",
                            "; ".join("rand %s: %s" % (x["rand"], x["mesaj"]) for x in er[:6])))
    with conn.cursor() as cur:
        cur.execute("DELETE FROM public.declaratii_depuse WHERE tenant_id=%s AND sursa='migrare'",
                    (tenant_id,))
        n = 0
        for r in randuri:
            cur.execute("""
                INSERT INTO public.declaratii_depuse (tenant_id, an, luna, tip, data_depunere, sursa)
                VALUES (%s,%s,%s,%s,%s,'migrare')
            """, (tenant_id, r.get("an"), r.get("luna"),
                  str(r["tip"]).strip().lower(),   # [tip_lowercase] canonic la STOCARE (CHECK-ul il impune)
                  r.get("data_depunere")))
            n += 1
    conn.commit()
    return {"importati": n}

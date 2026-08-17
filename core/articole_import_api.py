"""
core/articole_import_api.py - import articole + stoc initial CV (F151) per firma.
Citeste un export (.xlsx/.csv) cu denumire, UM, cantitate, pret unitar, cont stoc (optional).
Mapeaza flexibil coloanele (tiparul asociati_import_api). Creeaza articolele si o miscare
de INTRARE initiala (document 'sold initial') -> CMP-ul de pornire al fisei de magazie.
Idempotenta: articolele existente (aceeasi denumire) NU se dubleaza - se sare cu motiv.
"""
from __future__ import annotations


def _gaseste_col(antet, *chei):
    # 1) potrivire exacta (um != denUMire), 2) apoi continere
    for i, h in enumerate(antet):
        hl = str(h).strip().lower()
        if hl in chei:
            return i
    for i, h in enumerate(antet):
        hl = str(h).strip().lower()
        for k in chei:
            if k in hl:
                return i
    return -1


from core.numere import numar as _numar  # sursa unica (15.07.2026, vezi core/numere.py)


def extrage(continut, nume_fisier=""):
    """Intoarce [{denumire, um, cantitate, pret, cont_stoc, cont_cheltuiala, valid, motiv}]."""
    nume = (nume_fisier or "").lower()
    randuri = []
    if nume.endswith(".csv") or nume.endswith(".tsv") or nume.endswith(".txt"):
        import csv, io
        text = continut.decode("utf-8-sig", errors="replace") if isinstance(continut, bytes) else str(continut)
        prima = text.splitlines()[0] if text.splitlines() else ""
        delim = "\t" if nume.endswith(".tsv") else (";" if prima.count(";") > prima.count(",") else ",")
        randuri = list(csv.reader(io.StringIO(text), delimiter=delim))
    elif nume.endswith(".xlsx") or nume.endswith(".xlsm"):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(continut), read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            randuri.append(["" if c is None else c for c in r])
        wb.close()
    else:
        raise ValueError("format neacceptat (doar .csv sau .xlsx)")
    if not randuri:
        return []
    antet = [str(x) for x in randuri[0]]
    i_den = _gaseste_col(antet, "denumire", "articol", "produs", "material")
    i_um = _gaseste_col(antet, "um", "u.m", "unitate")
    i_cant = _gaseste_col(antet, "cantitate", "stoc", "cant")
    i_pret = _gaseste_col(antet, "pret", "cmp", "cost")
    i_cont = _gaseste_col(antet, "cont stoc", "cont_stoc")
    i_ch = _gaseste_col(antet, "cont cheltuiala", "cont_cheltuiala")
    if i_den < 0:
        raise ValueError("nu găsesc coloana cu denumirea articolului")
    rez = []
    for r in randuri[1:]:
        den = str(r[i_den] if i_den < len(r) else "").strip()
        if not den:
            continue
        cant = _numar(r[i_cant]) if 0 <= i_cant < len(r) else 0.0
        pret = _numar(r[i_pret]) if 0 <= i_pret < len(r) else 0.0
        um = (str(r[i_um]).strip() if 0 <= i_um < len(r) else "") or "buc"
        cont = (str(r[i_cont]).strip() if 0 <= i_cont < len(r) else "") or "302"
        ch = (str(r[i_ch]).strip() if 0 <= i_ch < len(r) else "") or "601"
        valid, motiv = True, "ok"
        if cant < 0 or pret < 0:
            valid, motiv = False, "cantitate/preț negativ"
        # [articol_pret] articol cu stoc dar pret 0/lipsa: parserul fabrica pret 0.0 (coloana absenta /
        # celula goala) -> miscarea de intrare ar avea valoare 0, CMP porneste de la 0, valoarea stocului
        # sub-raportata TACIT (DS cap.17, fara default fabricat). Un stoc real are cost > 0.
        elif cant > 0 and pret <= 0:
            valid, motiv = False, "are stoc dar preț unitar 0/lipsă: valoarea stocului ar fi 0 - completează costul unitar (CMP)"
        rez.append({"denumire": den[:255], "um": um[:20], "cantitate": cant, "pret": pret,
                    "cont_stoc": cont[:10], "cont_cheltuiala": ch[:10], "valid": valid, "motiv": motiv})
    return rez


def rezumat(articole):
    ok = [a for a in articole if a["valid"]]
    return {"total": len(articole), "valide": len(ok), "invalide": len(articole) - len(ok),
            "cu_stoc": sum(1 for a in ok if a["cantitate"] > 0),
            "valoare_totala": round(sum(a["cantitate"] * a["pret"] for a in ok), 2)}


def importa(conn, schema, articole, data_sold=None):
    """Creeaza articolele noi + intrarea initiala pentru cele cu stoc. Existentele se sar."""
    import datetime as _dt
    d = data_sold or _dt.date.today().isoformat()
    create, sarite = 0, []
    with conn.cursor() as cur:
        for a in articole:
            if not a["valid"]:
                sarite.append({"denumire": a["denumire"], "motiv": a["motiv"]})
                continue
            cur.execute(f"SELECT id FROM {schema}.articole WHERE lower(denumire)=lower(%s)", (a["denumire"],))
            if cur.fetchone():
                sarite.append({"denumire": a["denumire"], "motiv": "există deja"})
                continue
            cur.execute(f"""INSERT INTO {schema}.articole (denumire, um, cont_stoc, cont_cheltuiala)
                            VALUES (%s,%s,%s,%s) RETURNING id""",
                        (a["denumire"], a["um"], a["cont_stoc"], a["cont_cheltuiala"]))
            aid = cur.fetchone()[0]
            if a["cantitate"] > 0:
                val = round(a["cantitate"] * a["pret"], 2)
                cur.execute(f"""INSERT INTO {schema}.miscari_stoc (articol_id, data, tip, cantitate, pret_unitar, valoare, document)
                                VALUES (%s,%s,'intrare',%s,%s,%s,'sold initial (import)')""",
                            (aid, d, a["cantitate"], a["pret"], val))
            create += 1
    conn.commit()
    return {"create": create, "sarite": sarite}

"""
core/solduri_parteneri_api.py — import solduri parteneri (defalcare 4111/401 per client/furnizor).

Citeste un fisier (.xlsx/.csv) cu coloanele: Cont, CUI, Denumire, Sold debitor, Sold creditor.
Un rand = un partener pe un cont. Se salveaza in schema tenantului (tabel solduri_parteneri).

Coerenta cu stratul 2 (solduri_initiale): suma soldurilor partenerilor pe contul sintetic
(4111, 401) se compara cu soldul sintetic din balanta (analiticele 4111.X grupate pe radacina).
Informativ, NU blocant — contabilul decide.
"""
from __future__ import annotations


from core.numere import numar as _numar  # sursa unica (15.07.2026, vezi core/numere.py)
from core.migrare_api import respinge  # [P8/C] respingerea e o afirmatie, cu regula din nomenclator


def _gaseste_col(antet, *chei, exclus=None):
    """Indexul primei coloane al carei antet contine una din chei. -1 daca lipseste.
    `exclus` = indecsi deja atribuiti altui camp; nu pot fi si acesta.

    Fara `exclus`, un antet "CUI partener" se potriveste si la "cui", si la "partener":
    prima potrivire castiga, iar denumirea ajunge sa fie CUI-ul. Dovedit 15.07.2026
    prin migrare reala - cabinetele isi numesc chiar asa coloanele.
    """
    lua = set(x for x in (exclus or []) if x is not None) if not isinstance(exclus, int) else {exclus}
    for i, h in enumerate(antet):
        if i in lua:
            continue
        hl = str(h).strip().lower()
        for k in chei:
            if k in hl:
                return i
    return -1


def _radacina(cont):
    """Sinteticul de baza al unui cont: '4111.ALPHA' -> '4111', '401' -> '401'."""
    return str(cont).split(".")[0].strip()


def _curata_cui(v):
    """Normalizeaza CUI: RO14837428 -> 14837428, scoate spatii."""
    t = str(v or "").strip().upper().replace(" ", "")
    if t.startswith("RO"):
        t = t[2:]
    return t


def extrage(continut, nume_fisier=""):
    """
    Intoarce [{cont, cui, denumire, debit, credit}] din fisierul de parteneri.
    Detecteaza coloanele dupa antet (cont / cui / denumire / debit / credit).
    """
    nume = (nume_fisier or "").lower()
    randuri = []

    if nume.endswith(".csv") or nume.endswith(".tsv") or nume.endswith(".txt"):
        import csv, io
        text = continut.decode("utf-8-sig", errors="replace") if isinstance(continut, bytes) else str(continut)
        prima = text.splitlines()[0] if text.splitlines() else ""
        delim = "\t" if nume.endswith(".tsv") else (";" if prima.count(";") > prima.count(",") else ",")
        randuri = list(csv.reader(io.StringIO(text), delimiter=delim))
    elif nume.endswith(".xlsx") or nume.endswith(".xlsm"):
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(continut), read_only=True, data_only=True)
        ws = wb.active
        for r in ws.iter_rows(values_only=True):
            randuri.append(["" if c is None else c for c in r])
        wb.close()
    else:
        raise ValueError("format neacceptat (doar .csv sau .xlsx)")

    if not randuri:
        return []

    antet = [str(x) for x in randuri[0]]
    i_cont = _gaseste_col(antet, "cont", "simbol")
    i_cui = _gaseste_col(antet, "cui", "cif", "cod fiscal")
    # denumirea se cauta DUPA cont si cui, sarind peste coloanele lor
    i_den = _gaseste_col(antet, "denumire", "nume", "partener", exclus=[i_cont, i_cui])
    i_deb = _gaseste_col(antet, "debitor", "debit")
    i_cre = _gaseste_col(antet, "creditor", "credit")
    if i_cont < 0:
        i_cont = 0

    out = []
    for r in randuri[1:]:
        if i_cont >= len(r):
            continue
        cont = str(r[i_cont]).strip()
        if not cont:
            continue
        cui = _curata_cui(r[i_cui]) if (0 <= i_cui < len(r)) else ""
        den = str(r[i_den]).strip() if (0 <= i_den < len(r)) else ""
        deb = _numar(r[i_deb]) if (0 <= i_deb < len(r)) else 0.0
        cre = _numar(r[i_cre]) if (0 <= i_cre < len(r)) else 0.0
        out.append({"cont": cont, "cui": cui, "denumire": den, "debit": deb, "credit": cre})
    return out


def asigura_tabel(conn):
    """Creeaza tabelul solduri_parteneri in schema curenta (search_path = tenantul)."""
    with conn.cursor() as cur:
        cur.execute("""
            CREATE TABLE IF NOT EXISTS solduri_parteneri (
                id SERIAL PRIMARY KEY,
                cont TEXT NOT NULL,
                cui TEXT NOT NULL DEFAULT '',
                denumire TEXT NOT NULL DEFAULT '',
                sold_debitor NUMERIC(15,2) NOT NULL DEFAULT 0,
                sold_creditor NUMERIC(15,2) NOT NULL DEFAULT 0,
                data_referinta DATE
            )
        """)
    conn.commit()


def coerenta(conn, randuri):
    """
    Compara soldurile partenerilor pe sintetic (radacina) cu soldurile din solduri_initiale.
    Intoarce [{cont, suma_parteneri, sold_balanta, diferenta, coincide}] per sintetic.
    Daca solduri_initiale nu exista inca, sold_balanta = None (nu putem verifica).
    """
    # suma partenerilor pe radacina sintetica, sold net (debit - credit)
    pe_sintetic = {}
    for r in randuri:
        rad = _radacina(r["cont"])
        net = float(r.get("debit", 0) or 0) - float(r.get("credit", 0) or 0)
        pe_sintetic[rad] = pe_sintetic.get(rad, 0.0) + net

    # soldurile din balanta, grupate pe radacina
    bal = {}
    with conn.cursor() as cur:
        cur.execute("SELECT to_regclass('solduri_initiale')")
        exista = cur.fetchone()[0] is not None
        if exista:
            cur.execute("SELECT cont, sold_debitor, sold_creditor FROM solduri_initiale")
            for cont, sd, sc in cur.fetchall():
                rad = _radacina(cont)
                bal[rad] = bal.get(rad, 0.0) + (float(sd) - float(sc))

    out = []
    for rad in sorted(pe_sintetic.keys()):
        sp = round(pe_sintetic[rad], 2)
        if exista and rad in bal:
            sb = round(bal[rad], 2)
            dif = round(sp - sb, 2)
            out.append({"cont": rad, "suma_parteneri": sp, "sold_balanta": sb,
                        "diferenta": dif, "coincide": abs(dif) < 0.01})
        else:
            out.append({"cont": rad, "suma_parteneri": sp, "sold_balanta": None,
                        "diferenta": None, "coincide": None})
    return out


# Conturile pe care se tin solduri pe parteneri. Restul (5121 banca, 5311 casa,
# conturi de venituri/cheltuieli) NU au parteneri: un sold pe partener acolo e o
# eroare de export, nu o realitate contabila.
CONTURI_PARTENERI = ("4111", "401", "409", "419", "4118", "4091", "4092", "4093")

# cheia oficiala de control a CUI romanesc (aceeasi ca la CNP in salariati_import_api:
# algoritm publicat, nu inventat aici)
_CHEIE_CUI = [7, 5, 3, 2, 1, 7, 5, 3, 2]


def valideaza_cui(cui):
    """(valid, motiv) - cifra de control a CUI romanesc, LOCAL (fara retea).
    anaf_api.valideaza_cui() interogheaza ANAF si e potrivit la o firma, nu la fiecare
    rand dintr-un import de sute de parteneri."""
    c = "".join(ch for ch in str(cui or "") if ch.isdigit())
    if not c:
        return False, "lipsa"
    if not (2 <= len(c) <= 10):
        return False, "lungime (2-10 cifre)"
    corp, ctrl = c[:-1], int(c[-1])
    corp = corp.rjust(9, "0")
    s = sum(int(corp[i]) * _CHEIE_CUI[i] for i in range(9))
    rest = (s * 10) % 11
    if rest == 10:
        rest = 0
    return (rest == ctrl), ("ok" if rest == ctrl else "cifra de control")


def verifica_randuri(randuri):
    """PURA: (erori, randuri_bune). Ce nu poate intra in evidenta si de ce.

    Pana la 15.07.2026 nu exista NICIO validare intre extrage() si importa():
    - CUI cu cifra de control gresita -> intra
    - partener fara CUI pe cont de parteneri -> intra
    - partener pe 5121 (banca) -> intra
    verifica() calcula diferentele fata de balanta si le trimitea in raspuns, ecranul
    le afisa cu ⚠, iar importa() scria oricum. Migrarea marca stratul "gata" VERDE.
    Semnalarea fara oprire e mai rea decat tacerea: da impresia ca s-a verificat.
    """
    erori, bune = [], []
    for i, r in enumerate(randuri or [], start=2):   # +2: randul din fisier, cu antet
        cont = str(r.get("cont") or "").strip()
        cui = str(r.get("cui") or "").strip()
        den = str(r.get("denumire") or "").strip()
        rad = _radacina(cont)
        if rad not in CONTURI_PARTENERI:
            erori.append(respinge(
                "sold partener", i, "cont_nepartener",
                "contul %s nu ține solduri pe parteneri (doar %s)"
                % (cont or "?", ", ".join(CONTURI_PARTENERI[:4])), cont=cont))
            continue
        ok, motiv = valideaza_cui(cui)
        if not ok:
            erori.append(respinge(
                "sold partener", i, "cui_invalid",
                ("partenerul %s nu are CUI" % (den or "?")) if motiv == "lipsa"
                else "CUI %s invalid (%s)" % (cui, motiv),
                cont=cont, cui=cui, denumire=den))
            continue
        bune.append(r)
    return erori, bune


def importa(conn, randuri, data_referinta=None):
    """Inlocuieste partenerii (DELETE + INSERT). Intoarce {randuri, total_debit, total_credit}.
    Ridica ValueError daca vreun rand nu poate intra in evidenta (vezi verifica_randuri).
    Refuzul e PRIMA POARTA: nimic nu se scrie dintr-un import cu randuri invalide."""
    # [probare invalid, 03.09.2026 — aceeasi clasa ca importul de istoric] UN IMPORT GOL
    # STERGEA TOT, IN TACERE. Lista vida trecea de verificare (n-are ce respinge), ajungea la
    # `DELETE FROM solduri_parteneri` de mai jos, si raspundea ca un import reusit cu zero randuri. Un fisier gol,
    # unul cu numai antet (`extrage` intoarce [] pe el), sau un apel de API cu lista goala — toate
    # ajungeau acolo. *Un import care nu aduce nimic n-are voie sa stearga ce era acolo.*
    if not randuri:
        raise ValueError("nu ai trimis niciun rând. Importul ar fi șters soldurile inițiale ale "
                         "partenerilor și n-ar fi pus nimic în loc. Dacă chiar vrei să le "
                         "golești, e altă operațiune.")
    erori, _bune = verifica_randuri(randuri)
    if erori:
        det = "; ".join("rand %s: %s" % (e["rand"], e["mesaj"]) for e in erori[:6])
        if len(erori) > 6:
            det += " (și încă %d)" % (len(erori) - 6)
        raise ValueError(
            "%d rânduri nu pot intra în evidență: %s. Soldurile pe parteneri au nevoie "
            "de CUI valid (intră în D394 și SAF-T) și de un cont care ține parteneri."
            % (len(erori), det))
    asigura_tabel(conn)
    td = tc = 0.0
    with conn.cursor() as cur:
        cur.execute("DELETE FROM solduri_parteneri")
        for r in randuri:
            cur.execute(
                "INSERT INTO solduri_parteneri (cont, cui, denumire, sold_debitor, sold_creditor, data_referinta) "
                "VALUES (%s,%s,%s,%s,%s,%s)",
                (r["cont"], r.get("cui", ""), r.get("denumire", ""),
                 r.get("debit", 0), r.get("credit", 0), data_referinta))
            td += float(r.get("debit", 0) or 0)
            tc += float(r.get("credit", 0) or 0)
    conn.commit()
    return {"randuri": len(randuri), "total_debit": round(td, 2), "total_credit": round(tc, 2)}


def rezumat(conn):
    """{are_parteneri, randuri, total_debit, total_credit} pentru firma curenta."""
    asigura_tabel(conn)
    with conn.cursor() as cur:
        cur.execute("SELECT count(*), COALESCE(sum(sold_debitor),0), COALESCE(sum(sold_creditor),0) FROM solduri_parteneri")
        n, td, tc = cur.fetchone()
    return {"are_parteneri": n > 0, "randuri": n, "total_debit": float(td), "total_credit": float(tc)}
